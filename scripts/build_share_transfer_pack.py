"""
BitWealth Asset Managers (Pty) Ltd — Founding Shares Transfer Pack
Davin -> Mhuri Investment Holdings (Pty) Ltd (100 Ordinary Shares / 10%)
Gives effect to SHA clause 7 and Deed of Adherence clause 2.2.

Contains: (1) Directors' Resolution, (2) Instrument of Transfer,
(3) Updated Securities Register extract, (4) Remaining-steps checklist.

Output: docs/Shareholding/BitWealth_Share_Transfer_Pack_Mhuri_v1.docx

NOT legal/tax advice — the share valuation and STT figures are placeholders
pending your tax advisor's confirmation of current fair value.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR = Path("docs") / "Shareholding"
OUT_DIR.mkdir(parents=True, exist_ok=True)

FOUNDER    = "Davin Harald Gaier"
FOUNDER_ID = "8405025239081"
PARTNER    = "Simon Henry Newbold Hobday"
CO_NAME    = "BitWealth Asset Managers (Pty) Ltd"
CO_REG     = "2026/090346/07"
SIMCO_NAME = "Mhuri Investment Holdings (Pty) Ltd"
SIMCO_REG  = "2017/334996/07"
SIMCO_ADDR = "6 Sardinia, Golf Close, Lonehill, Gauteng, 2196"

NAVY = RGBColor(0x0A, 0x2A, 0x43)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)
RED  = RGBColor(0xB7, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY_HEX = "0A2A43"

BODY_FONT = "Calibri"
BODY_SZ = 11


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def para(doc, text="", bold=False, italic=False, size=BODY_SZ, color=None,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, sb=2, sa=4, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after = Pt(sa)
    p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    if text:
        run = p.add_run(text)
        run.bold = bold; run.italic = italic
        run.font.name = BODY_FONT; run.font.size = Pt(size)
        run.font.color.rgb = color if color else DARK
    return p


def clause(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rn = p.add_run(f"{number}\t"); rn.bold = True
    rn.font.name = BODY_FONT; rn.font.size = Pt(BODY_SZ); rn.font.color.rgb = DARK
    rt = p.add_run(text)
    rt.font.name = BODY_FONT; rt.font.size = Pt(BODY_SZ); rt.font.color.rgb = DARK
    return p


def section(doc, title, page_break=True, sb=4):
    if page_break:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title.upper()); r.bold = True
    r.font.name = BODY_FONT; r.font.size = Pt(15); r.font.color.rgb = NAVY
    return p


def hrule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1"); b.set(qn("w:color"), NAVY_HEX)
    pBdr.append(b); pPr.append(pBdr)


def sig_line(doc, label, name, capacity):
    for line in [
        "Signed: __________________________\t\tDate: __________________________",
        f"Full Name: {name}",
        f"Capacity: {capacity}",
    ]:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_before = Pt(1); p2.paragraph_format.space_after = Pt(3)
        r2 = p2.add_run(line); r2.font.name = BODY_FONT; r2.font.size = Pt(BODY_SZ)
        r2.font.color.rgb = DARK
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


doc = Document()
style = doc.styles["Normal"]
style.font.name = BODY_FONT; style.font.size = Pt(BODY_SZ)
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.5)

# ── Cover ─────────────────────────────────────────────────────────────────
para(doc, "FOUNDING SHARES TRANSFER PACK", bold=True, size=20, color=NAVY,
     align=WD_ALIGN_PARAGRAPH.CENTER, sb=0, sa=6)
para(doc, CO_NAME + f" (Registration No. {CO_REG})", bold=True, size=13, color=NAVY,
     align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)
para(doc, "Transfer of 100 Ordinary Shares (10%) from Davin Harald Gaier to "
          "Mhuri Investment Holdings (Pty) Ltd", size=11, color=GREY,
     align=WD_ALIGN_PARAGRAPH.CENTER, sa=10)
hrule(doc)
para(doc,
    "WORKING DRAFT — NOT REVIEWED BY AN ATTORNEY OR TAX ADVISOR. The share valuation and "
    "STT figures below are placeholders pending confirmation of current fair value by your "
    "tax advisor. This pack gives effect to SHA clause 7 and Deed of Adherence clause 2.2 "
    "and must be finalised before signature.",
    size=9, italic=True, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER, sa=10)
hrule(doc)

# ══════════════════════════════════════════════════════════════════════════
# 1. DIRECTORS' RESOLUTION
# ══════════════════════════════════════════════════════════════════════════
section(doc, "1. Written Resolution of Directors", page_break=False, sb=16)
para(doc, f"{CO_NAME} (Registration No. {CO_REG}) (\u201cthe Company\u201d)", bold=True, sa=8)
para(doc, "WRITTEN RESOLUTION OF THE DIRECTORS passed in terms of section 74 of the "
          "Companies Act 71 of 2008 and the Company's Memorandum of Incorporation, "
          "on [insert date].")

clause(doc, "1.", "CONDITIONS PRECEDENT. The directors record that the conditions "
       "precedent in clause 7.4 of the Shareholders\u2019 Agreement dated [insert SHA date] "
       "and clause 2.2 of the Deed of Adherence dated [insert Deed date] have been "
       "satisfied, namely: (a) Simon Henry Newbold Hobday has obtained independent legal "
       "advice; (b) independent tax advice has been obtained regarding the income tax, "
       "CGT and donations tax consequences of the transfer, including the consequences of "
       "transferring to Mhuri Investment Holdings (Pty) Ltd rather than to Simon "
       "personally; (c) Finova (Pty) Ltd (FSP No. 21095) has provided a Juristic "
       "Representative appointment or letter of intent to appoint; and (d) Simon meets the "
       "minimum FAIS fit-and-proper requirements applicable to his representative role.")
clause(doc, "2.", "APPROVAL OF TRANSFER. The directors approve the transfer of 100 "
       "(one hundred) Ordinary Shares, constituting 10% of the issued share capital of the "
       "Company, from Davin Harald Gaier to Mhuri Investment Holdings (Pty) Ltd "
       "(Registration No. 2017/334996/07), for a consideration of [R_______ / nil, subject "
       "to fair value determination for Securities Transfer Tax purposes].")
clause(doc, "3.", "SECURITIES REGISTER. The directors authorise and direct that the "
       "Company\u2019s securities register be updated to reflect Davin Harald Gaier as holder "
       "of 900 Ordinary Shares and Mhuri Investment Holdings (Pty) Ltd as holder of 100 "
       "Ordinary Shares, and that any share certificates be issued or cancelled "
       "accordingly.")
clause(doc, "4.", "SECURITIES TRANSFER TAX. The directors authorise the Company to settle "
       "the Securities Transfer Tax payable on the transfer, as required by clause 7.5 of "
       "the Shareholders\u2019 Agreement, and to file the relevant SARS return.")
clause(doc, "5.", "BENEFICIAL OWNERSHIP FILING. The directors authorise the Company to "
       "update its Beneficial Ownership register and to file the required disclosure with "
       "the Companies and Intellectual Property Commission, recording Simon Henry Newbold "
       "Hobday and his spouse as the natural persons with a beneficial interest in the "
       "shares held by Mhuri Investment Holdings (Pty) Ltd.")
clause(doc, "6.", "GENERAL AUTHORITY. Any one director of the Company is authorised to "
       "sign all documents and take all steps necessary or incidental to give effect to "
       "this resolution.")

doc.add_paragraph()
sig_line(doc, "Director", FOUNDER, "Managing Director / Founder")

hrule(doc)

# ══════════════════════════════════════════════════════════════════════════
# 2. INSTRUMENT OF TRANSFER
# ══════════════════════════════════════════════════════════════════════════
section(doc, "2. Instrument of Transfer")
para(doc, CO_NAME + f" (Registration No. {CO_REG})", bold=True, sa=8)
para(doc, "This Instrument of Transfer is made in terms of the Shareholders\u2019 Agreement "
          "dated [insert SHA date] and the Deed of Adherence dated [insert Deed date].")

t = doc.add_table(rows=6, cols=2)
t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
rows = [
    ("Transferor", f"{FOUNDER} (ID {FOUNDER_ID})"),
    ("Transferee", f"{SIMCO_NAME} (Reg. No. {SIMCO_REG})\n{SIMCO_ADDR}"),
    ("Class of shares", "Ordinary Shares of R1.00 par value each"),
    ("Number of shares transferred", "100 (one hundred), being 10% of the issued share capital"),
    ("Consideration", "[R_______ / nil consideration, subject to fair value for STT purposes]"),
    ("Date of transfer", "[insert date]"),
]
for i, (k, v) in enumerate(rows):
    c0 = t.cell(i, 0); c0.text = ""
    r0 = c0.paragraphs[0].add_run(k); r0.bold = True
    r0.font.name = BODY_FONT; r0.font.size = Pt(BODY_SZ); r0.font.color.rgb = WHITE
    set_cell_bg(c0, NAVY_HEX)
    c1 = t.cell(i, 1); c1.text = ""
    r1 = c1.paragraphs[0].add_run(v)
    r1.font.name = BODY_FONT; r1.font.size = Pt(BODY_SZ); r1.font.color.rgb = DARK

doc.add_paragraph()
para(doc, "The Transferor hereby transfers, and the Transferee hereby accepts transfer of, "
          "the shares described above, to be held subject to the Shareholders\u2019 Agreement "
          "and the Deed of Adherence referred to above.", sb=8)

doc.add_paragraph()
sig_line(doc, "Transferor", FOUNDER, "In his personal capacity")
sig_line(doc, "Transferee", "[Authorised director]", f"For and on behalf of {SIMCO_NAME}")

hrule(doc)

# ══════════════════════════════════════════════════════════════════════════
# 3. UPDATED SECURITIES REGISTER EXTRACT
# ══════════════════════════════════════════════════════════════════════════
section(doc, "3. Updated Securities Register Extract")
para(doc, "Extract from the Company\u2019s securities register (maintained under section 50 "
          "of the Companies Act 71 of 2008) following registration of the transfer:", sa=8)

t2 = doc.add_table(rows=4, cols=4)
t2.style = "Table Grid"; t2.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["Shareholder", "Ordinary Shares", "% Holding", "Date registered"]):
    c = t2.cell(0, i); c.text = ""
    r = c.paragraphs[0].add_run(h); r.bold = True
    r.font.name = BODY_FONT; r.font.size = Pt(10); r.font.color.rgb = WHITE
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cell_bg(c, NAVY_HEX)
data_rows = [
    (FOUNDER, "900", "90%", "—"),
    (SIMCO_NAME, "100", "10%", "[insert date]"),
    ("TOTAL", "1,000", "100%", ""),
]
for r_idx, row in enumerate(data_rows, 1):
    bold = row[0] == "TOTAL"
    for c_idx, val in enumerate(row):
        c = t2.cell(r_idx, c_idx); c.text = ""
        run = c.paragraphs[0].add_run(val)
        run.bold = bold
        run.font.name = BODY_FONT; run.font.size = Pt(10); run.font.color.rgb = DARK
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

hrule(doc)

# ══════════════════════════════════════════════════════════════════════════
# 4. REMAINING STEPS CHECKLIST
# ══════════════════════════════════════════════════════════════════════════
section(doc, "4. Remaining Steps Checklist")
checklist = [
    "Tax advisor confirms current fair value of the 100 shares and the STT amount payable.",
    "Directors' resolution (Section 1) signed and dated.",
    "Instrument of Transfer (Section 2) signed by Davin and an authorised Mhuri director.",
    "Securities Transfer Tax return filed and paid via SARS eFiling (within 2 months of transfer).",
    "Company's internal securities register updated (Section 3) and share certificates issued/cancelled.",
    "Beneficial Ownership register updated and filed with CIPC, recording Simon and his "
    "spouse as beneficial owners behind Mhuri.",
    "Mhuri submits a Dividends Tax exemption declaration to the Company before any dividend "
    "is paid.",
    "FICA/KYC documents obtained for Mhuri (registration documents, registered address "
    "proof, directors' ID and proof of address).",
    "Finova's compliance officer informed of the change in BitWealth's shareholding.",
    "Signed SHA, Deed of Adherence (with spousal consent), resolution and transfer form "
    "filed together in the Company's minute book.",
]
for item in checklist:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("\u2610\t"); r.font.name = BODY_FONT; r.font.size = Pt(BODY_SZ)
    r2 = p.add_run(item); r2.font.name = BODY_FONT; r2.font.size = Pt(BODY_SZ); r2.font.color.rgb = DARK

path = OUT_DIR / "BitWealth_Share_Transfer_Pack_Mhuri_v1.docx"
doc.save(path)
print(f"Saved: {path}")
